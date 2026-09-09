"""
apps/presupuestos/xlsx_exports.py — Generadores de xlsx para Despiece y APU.

Funciones principales:
    build_despiece_xlsx(proyecto)          → BytesIO (hoja Despiece — modelo legacy DespieceLinea)
    build_despiece_maestro_xlsx(despiece)  → BytesIO (hoja Despiece — DespieceMaestroLinea)
    build_apu_xlsx(apu)                    → BytesIO con el workbook (2 hojas)

Requiere: openpyxl (pip install openpyxl)
"""

from __future__ import annotations

import io
from decimal import Decimal
from typing import TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment, Border, Font, GradientFill, PatternFill, Side
)
from openpyxl.utils import get_column_letter

if TYPE_CHECKING:
    from apps.presupuestos.models import APUProyecto
    from apps.comercial.models import Proyecto


# ── Paleta ───────────────────────────────────────────────────────────────────

_BRAND_BLUE    = "FF1470E6"
_BRAND_DARK    = "FF0F172A"
_SECTION_BG    = "FFE8F0FD"   # azul muy claro para cabeceras de sección
_HEADER_BG     = "FF1470E6"   # azul marca para la cabecera principal
_SUBTOTAL_BG   = "FFD1E3FC"   # fila subtotal
_ALT_ROW       = "FFF7F9FF"   # filas alternas
_BORDER_GRAY   = "FFCBD5E1"

_WHITE = "FFFFFFFF"
_DARK  = "FF0F172A"


# ── Helpers de estilos ────────────────────────────────────────────────────────

def _font(bold=False, size=10, color=_DARK, name="Arial"):
    return Font(name=name, size=size, bold=bold, color=color)


def _fill(hex_color: str):
    return PatternFill("solid", fgColor=hex_color)


def _border(color=_BORDER_GRAY):
    side = Side(style="thin", color=color)
    return Border(left=side, right=side, top=side, bottom=side)


def _align(horizontal="left", vertical="center", wrap=False):
    return Alignment(horizontal=horizontal, vertical=vertical, wrap_text=wrap)


def _money(value) -> str:
    """Formatea a número entero con separador de miles."""
    try:
        return f"{int(round(float(value))):,}"
    except (TypeError, ValueError):
        return "—"


def _qty(value) -> str:
    """Formatea cantidad con hasta 4 decimales eliminando ceros finales."""
    try:
        f = float(value)
        if f == int(f):
            return str(int(f))
        return f"{f:.4f}".rstrip("0")
    except (TypeError, ValueError):
        return "—"


def _write_row(ws, row_num: int, values: list, fonts=None, fills=None,
               aligns=None, border=True, height=None):
    """Escribe una fila con estilos opcionales por columna."""
    for col_idx, val in enumerate(values, start=1):
        cell = ws.cell(row=row_num, column=col_idx, value=val)
        if fonts and col_idx - 1 < len(fonts) and fonts[col_idx - 1]:
            cell.font = fonts[col_idx - 1]
        if fills and col_idx - 1 < len(fills) and fills[col_idx - 1]:
            cell.fill = fills[col_idx - 1]
        if aligns and col_idx - 1 < len(aligns) and aligns[col_idx - 1]:
            cell.alignment = aligns[col_idx - 1]
        if border:
            cell.border = _border()
    if height:
        ws.row_dimensions[row_num].height = height


def _merge_write(ws, row, col_start, col_end, value, font=None, fill=None,
                 align=None, border=True, height=None):
    """Escribe un valor en celdas fusionadas."""
    ws.merge_cells(
        start_row=row, start_column=col_start,
        end_row=row, end_column=col_end,
    )
    cell = ws.cell(row=row, column=col_start, value=value)
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if align:
        cell.alignment = align
    if border:
        cell.border = _border()
    if height:
        ws.row_dimensions[row].height = height
    return cell


# ════════════════════════════════════════════════════════════════════════════
# 1. HOJA DESPIECE
# ════════════════════════════════════════════════════════════════════════════

_DESPIECE_COLS = [
    ("Producto asignado",    40),
    ("Proveedor",            28),
    ("Cantidad",             12),
    ("Unidad",               10),
    ("Precio Unit. (COP)",   18),
    ("Subtotal (COP)",       18),
]


def _build_despiece_sheet(ws, proyecto):
    """
    Rellena la hoja `ws` con las líneas de despiece del proyecto.
    Columnas: Producto, Proveedor, Cantidad, Unidad, Precio Unit., Subtotal.
    """
    # ── Anchos de columna ─────────────────────────────────────────────────
    for idx, (_, width) in enumerate(_DESPIECE_COLS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    num_cols = len(_DESPIECE_COLS)
    row = 1

    # ── Fila título ───────────────────────────────────────────────────────
    _merge_write(
        ws, row, 1, num_cols,
        f"LISTADO DE MATERIALES — {proyecto.nombre or proyecto.consecutivo}",
        font=_font(bold=True, size=13, color=_WHITE),
        fill=_fill(_BRAND_BLUE),
        align=_align("center"),
        height=28,
    )
    row += 1

    # ── Fila proyecto ─────────────────────────────────────────────────────
    _merge_write(
        ws, row, 1, 3,
        f"Proyecto: {proyecto.consecutivo}",
        font=_font(bold=True, size=10),
        fill=_fill("FFF0F4FF"),
        align=_align("left"),
        height=20,
    )
    _merge_write(
        ws, row, 4, num_cols,
        f"Cliente: {proyecto.cliente.razon_social if proyecto.cliente_id else '—'}",
        font=_font(size=10),
        fill=_fill("FFF0F4FF"),
        align=_align("left"),
    )
    row += 1

    # ── Cabecera de columnas ───────────────────────────────────────────────
    header_font  = _font(bold=True, size=10, color=_WHITE)
    header_fill  = _fill(_BRAND_DARK)
    header_align = _align("center")
    col_names = [c[0] for c in _DESPIECE_COLS]
    _write_row(ws, row, col_names,
               fonts=[header_font] * num_cols,
               fills=[header_fill] * num_cols,
               aligns=[header_align] * num_cols,
               height=20)
    row += 1

    # ── Cargar líneas ─────────────────────────────────────────────────────
    from apps.presupuestos.models import DespieceLinea

    lineas = (
        DespieceLinea.objects
        .filter(proyecto=proyecto)
        .select_related(
            "proyecto_sistema__sistema",
            "proyecto_sistema__subsistema",
            "producto",
        )
        .prefetch_related(
            "producto__proveedores_producto__proveedor",
        )
        .order_by(
            "proyecto_sistema__sistema__nombre",
            "proyecto_sistema__subsistema__nombre",
            "componente_codigo",
        )
    )

    total_general = Decimal("0")
    alt = False

    for linea in lineas:
        producto   = linea.producto.nombre if linea.producto_id else "Sin asignar"
        # Proveedor: el de menor precio activo (mismo que usa capturar_precio)
        proveedor  = "—"
        if linea.producto_id:
            pp = (
                linea.producto.proveedores_producto
                .filter(activo=True)
                .order_by("precio_unitario")
                .first()
            )
            if pp:
                proveedor = pp.proveedor.nombre
        cantidad   = linea.cantidad_final or Decimal("0")
        precio     = linea.precio_snapshot or Decimal("0")
        subtotal   = Decimal(str(cantidad)) * Decimal(str(precio))
        total_general += subtotal

        row_fill = _fill(_ALT_ROW) if alt else None
        alt = not alt

        num_style = _align("right")
        _write_row(
            ws, row,
            [producto, proveedor, _qty(cantidad), "und", _money(precio), _money(subtotal)],
            fonts=[_font(size=9.5)] * num_cols,
            fills=[row_fill] * num_cols,
            aligns=[
                _align(), _align(),
                num_style, _align("center"), num_style, num_style,
            ],
            height=16,
        )
        row += 1

    # ── Fila total ─────────────────────────────────────────────────────────
    total_fill = _fill(_SUBTOTAL_BG)
    _merge_write(
        ws, row, 1, 5,
        "TOTAL GENERAL",
        font=_font(bold=True, size=10),
        fill=total_fill,
        align=_align("right"),
        height=20,
    )
    cell = ws.cell(row=row, column=6, value=_money(total_general))
    cell.font  = _font(bold=True, size=10)
    cell.fill  = total_fill
    cell.alignment = _align("right")
    cell.border    = _border()

    # Fijar la fila de encabezado de columnas (row 3)
    ws.freeze_panes = "A4"


# ════════════════════════════════════════════════════════════════════════════
# 2. HOJA APU
# ════════════════════════════════════════════════════════════════════════════

_APU_COLS = [
    ("Descripción",         40),
    ("Rendimiento",         14),
    ("Unidad",              12),
    ("Precio Ref. (COP)",   18),
    ("Costo Unit. (COP)",   18),
    ("Costo Total (COP)",   18),
    ("Valor Unit. (COP)",   18),
    ("Valor Total (COP)",   18),
]

_TIPO_INFO = [
    ("MATERIALES",          "1. Materiales",           "FF1470E6"),
    ("HERRAMIENTAS_EQUIPOS","2. Herramientas y Equipos","FFe07d10"),
    ("TRANSPORTE",          "3. Transporte",            "FF7048d0"),
    ("MANO_DE_OBRA",        "4. Mano de Obra",          "FF17a85e"),
    ("ADMINISTRACION",      "5. Administración",        "FF64748b"),
]


def _build_apu_sheet(ws, apu):
    """
    Rellena la hoja `ws` con los datos del APU.
    """
    num_cols = len(_APU_COLS)

    for idx, (_, width) in enumerate(_APU_COLS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    row = 1

    # ── Título principal ──────────────────────────────────────────────────
    _merge_write(
        ws, row, 1, num_cols,
        "ANÁLISIS DE PRECIOS UNITARIOS",
        font=_font(bold=True, size=14, color=_WHITE),
        fill=_fill(_BRAND_BLUE),
        align=_align("center"),
        height=30,
    )
    row += 1

    # ── Nombre del APU ────────────────────────────────────────────────────
    _merge_write(
        ws, row, 1, num_cols,
        apu.nombre,
        font=_font(bold=True, size=11),
        fill=_fill("FFF0F4FF"),
        align=_align("left"),
        height=22,
    )
    row += 1

    # ── Datos del proyecto ────────────────────────────────────────────────
    ps      = apu.proyecto_sistema
    proyecto = ps.proyecto if ps else None
    sistema   = ps.sistema.nombre   if ps and ps.sistema_id   else "—"
    subsistema = ps.subsistema.nombre if ps and ps.subsistema_id else "—"
    cliente   = proyecto.cliente.razon_social if (proyecto and proyecto.cliente_id) else "—"
    consec    = proyecto.consecutivo if proyecto else "—"

    info_rows = [
        (f"Proyecto: {consec}", f"Cliente: {cliente}"),
        (f"Sistema: {sistema}", f"Subsistema: {subsistema}"),
        (
            f"Margen material: {apu.margen_material_pct}%    "
            f"IVA: {apu.iva_pct}%    "
            f"AIU contratista: {apu.aiu_contratista_pct}%    "
            f"Margen mano de obra: {apu.margen_mano_obra_pct}%",
            f"Días de duración: {apu.dias_duracion}"
        ),
    ]
    info_fill = _fill("FFF8FAFF")
    for left, right in info_rows:
        _merge_write(ws, row, 1, 4, left,
                     font=_font(size=9.5), fill=info_fill, align=_align("left"), height=17)
        _merge_write(ws, row, 5, num_cols, right,
                     font=_font(size=9.5), fill=info_fill, align=_align("left"))
        row += 1

    row += 1  # separador

    # ── Cargar todas las líneas ────────────────────────────────────────────
    from apps.common.choices import TipoAPU

    all_lineas = list(
        apu.lineas.select_related("item_catalogo__categoria", "despiece_linea__producto")
        .order_by("tipo", "descripcion")
    )

    totales_costo = Decimal("0")
    totales_valor = Decimal("0")

    for tipo_key, tipo_label, tipo_color in _TIPO_INFO:
        lineas_tipo = [l for l in all_lineas if l.tipo == tipo_key]

        # ── Cabecera de sección ───────────────────────────────────────────
        _merge_write(
            ws, row, 1, num_cols,
            tipo_label,
            font=_font(bold=True, size=10, color=_WHITE),
            fill=_fill(tipo_color),
            align=_align("left"),
            height=20,
        )
        row += 1

        if not lineas_tipo:
            _merge_write(
                ws, row, 1, num_cols,
                "(sin líneas registradas)",
                font=_font(size=9, color="FF94A3B8"),
                align=_align("center"),
                height=15,
                border=False,
            )
            row += 1
            continue

        # ── Cabecera de columnas ──────────────────────────────────────────
        col_header_fill = _fill(_SECTION_BG)
        _write_row(
            ws, row,
            [c[0] for c in _APU_COLS],
            fonts=[_font(bold=True, size=9.5, color=_DARK)] * num_cols,
            fills=[col_header_fill] * num_cols,
            aligns=[_align("center")] * num_cols,
            height=18,
        )
        row += 1

        # ── Filas de detalle ──────────────────────────────────────────────
        subtotal_costo = Decimal("0")
        subtotal_valor = Decimal("0")
        alt = False

        for linea in lineas_tipo:
            # Para materiales, resolver nombre desde despiece_linea si existe
            descripcion = linea.descripcion
            if linea.despiece_linea and linea.despiece_linea.producto:
                descripcion = linea.despiece_linea.producto.nombre

            precio_ref = linea.precio_referencia or Decimal("0")
            costo_u    = linea.costo_unitario    or Decimal("0")
            costo_t    = linea.costo_total       or Decimal("0")
            valor_u    = linea.valor_unitario    or Decimal("0")
            valor_t    = linea.valor_total       or Decimal("0")
            rend       = linea.rendimiento       or Decimal("1")

            subtotal_costo += costo_t
            subtotal_valor += valor_t

            row_fill = _fill(_ALT_ROW) if alt else None
            alt = not alt
            num_a = _align("right")

            _write_row(
                ws, row,
                [descripcion, _qty(rend), linea.get_unidad_display(),
                 _money(precio_ref), _money(costo_u), _money(costo_t),
                 _money(valor_u), _money(valor_t)],
                fonts=[_font(size=9.5)] * num_cols,
                fills=[row_fill] * num_cols,
                aligns=[_align("left", wrap=True),
                        _align("right"), _align("center"),
                        num_a, num_a, num_a, num_a, num_a],
                height=15,
            )
            row += 1

        # ── Subtotal de sección ───────────────────────────────────────────
        totales_costo += subtotal_costo
        totales_valor += subtotal_valor

        sub_fill = _fill(_SUBTOTAL_BG)
        _merge_write(
            ws, row, 1, 5,
            f"Subtotal {tipo_label}",
            font=_font(bold=True, size=9.5),
            fill=sub_fill,
            align=_align("right"),
            height=18,
        )
        ws.cell(row=row, column=6, value=_money(subtotal_costo)).font = _font(bold=True, size=9.5)
        ws.cell(row=row, column=6).fill  = sub_fill
        ws.cell(row=row, column=6).alignment = _align("right")
        ws.cell(row=row, column=6).border = _border()

        ws.cell(row=row, column=7, value="").fill = sub_fill
        ws.cell(row=row, column=7).border = _border()

        ws.cell(row=row, column=8, value=_money(subtotal_valor)).font = _font(bold=True, size=9.5)
        ws.cell(row=row, column=8).fill  = sub_fill
        ws.cell(row=row, column=8).alignment = _align("right")
        ws.cell(row=row, column=8).border = _border()

        row += 2  # salto entre secciones

    # ── Total general ─────────────────────────────────────────────────────
    total_fill = _fill("FFD1E3FC")
    _merge_write(
        ws, row, 1, 5,
        "TOTAL COSTOS DIRECTOS",
        font=_font(bold=True, size=10, color=_BRAND_DARK),
        fill=total_fill,
        align=_align("right"),
        height=22,
    )
    ws.cell(row=row, column=6, value=_money(totales_costo)).font = _font(bold=True, size=10, color=_BRAND_DARK)
    ws.cell(row=row, column=6).fill  = total_fill
    ws.cell(row=row, column=6).alignment = _align("right")
    ws.cell(row=row, column=6).border = _border()

    ws.cell(row=row, column=7, value="").fill = total_fill
    ws.cell(row=row, column=7).border = _border()

    ws.cell(row=row, column=8, value=_money(totales_valor)).font = _font(bold=True, size=10, color=_BRAND_DARK)
    ws.cell(row=row, column=8).fill  = total_fill
    ws.cell(row=row, column=8).alignment = _align("right")
    ws.cell(row=row, column=8).border = _border()

    # Fijar las primeras 6 filas de cabecera
    ws.freeze_panes = "A7"


# ════════════════════════════════════════════════════════════════════════════
# 2.5  HOJA DESPIECE MAESTRO  (DespieceMaestro → DespieceMaestroLinea)
# ════════════════════════════════════════════════════════════════════════════

_DESPIECE_MAESTRO_COLS = [
    ("Subconjunto",          20),
    ("Código",               12),
    ("Componente",           28),
    ("Categoría",            18),
    ("Producto asignado",    32),
    ("Cantidad",             10),
    ("Unidad",                8),
    ("Precio Unit. (COP)",   18),
    ("Subtotal (COP)",       18),
]


def _build_despiece_maestro_sheet(ws, despiece):
    """
    Rellena la hoja `ws` con las líneas de un DespieceMaestro.
    Lee DespieceMaestroLinea directamente (no DespieceLinea del modelo legacy).
    """
    for idx, (_, width) in enumerate(_DESPIECE_MAESTRO_COLS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    num_cols = len(_DESPIECE_MAESTRO_COLS)
    row = 1

    # ── Fila título ───────────────────────────────────────────────────────
    titulo = (
        f"LISTADO DE MATERIALES — "
        f"{despiece.nombre or despiece.subsistema.nombre}"
    )
    _merge_write(
        ws, row, 1, num_cols, titulo,
        font=_font(bold=True, size=13, color=_WHITE),
        fill=_fill(_BRAND_BLUE),
        align=_align("center"),
        height=28,
    )
    row += 1

    # ── Fila subsistema / sistema ─────────────────────────────────────────
    subsistema = despiece.subsistema
    sistema_nombre = subsistema.sistema.nombre if hasattr(subsistema, "sistema") else "—"
    _merge_write(
        ws, row, 1, 4,
        f"Subsistema: {subsistema.codigo} — {subsistema.nombre}",
        font=_font(bold=True, size=10),
        fill=_fill("FFF0F4FF"),
        align=_align("left"),
        height=20,
    )
    _merge_write(
        ws, row, 5, num_cols,
        f"Sistema: {sistema_nombre}",
        font=_font(size=10),
        fill=_fill("FFF0F4FF"),
        align=_align("left"),
    )
    row += 1

    # ── Fila proyecto / cliente ────────────────────────────────────────────
    proyecto = despiece.proyecto
    proyecto_texto = proyecto.consecutivo if proyecto else "—"
    cliente_texto = (
        proyecto.cliente.razon_social
        if (proyecto and proyecto.cliente_id) else "—"
    )
    _merge_write(
        ws, row, 1, 4,
        f"Proyecto: {proyecto_texto}",
        font=_font(bold=True, size=10),
        fill=_fill("FFF0F4FF"),
        align=_align("left"),
        height=20,
    )
    _merge_write(
        ws, row, 5, num_cols,
        f"Cliente: {cliente_texto}",
        font=_font(size=10),
        fill=_fill("FFF0F4FF"),
        align=_align("left"),
    )
    row += 1

    # ── Cabecera de columnas ───────────────────────────────────────────────
    col_names = [c[0] for c in _DESPIECE_MAESTRO_COLS]
    _write_row(
        ws, row, col_names,
        fonts=[_font(bold=True, size=10, color=_WHITE)] * num_cols,
        fills=[_fill(_BRAND_DARK)] * num_cols,
        aligns=[_align("center")] * num_cols,
        height=20,
    )
    row += 1

    # Fijar cabecera (4 filas de encabezado)
    ws.freeze_panes = "A5"

    # ── Líneas ────────────────────────────────────────────────────────────
    lineas = despiece.lineas.order_by("orden")

    total_general = Decimal("0")
    alt = False
    current_subconjunto = None

    col_aligns = [
        _align("left"),    # Subconjunto
        _align("left"),    # Código
        _align("left"),    # Componente
        _align("left"),    # Categoría
        _align("left"),    # Producto asignado
        _align("right"),   # Cantidad
        _align("center"),  # Unidad
        _align("right"),   # Precio Unit.
        _align("right"),   # Subtotal
    ]

    for linea in lineas:
        subconjunto_nom = linea.subconjunto_nombre or "General"

        # Fila separadora cuando cambia el subconjunto
        if subconjunto_nom != current_subconjunto:
            current_subconjunto = subconjunto_nom
            alt = False  # reinicia alternancia en cada sección
            _merge_write(
                ws, row, 1, num_cols,
                f"▶  {subconjunto_nom}",
                font=_font(bold=True, size=10),
                fill=_fill(_SECTION_BG),
                align=_align("left"),
                height=18,
            )
            row += 1

        fill_row = _fill(_ALT_ROW) if alt else None
        alt = not alt

        qty = (
            linea.cantidad_redondeada
            if linea.cantidad_redondeada is not None
            else linea.cantidad_calculada
        )
        precio_u = linea.precio_unitario
        precio_t = linea.precio_total

        values = [
            "",                                      # col Subconjunto (ya tiene sección)
            linea.componente_codigo,
            linea.componente_nombre,
            linea.categoria_nombre,
            linea.producto_nombre or "—",
            _qty(qty),
            linea.unidad,
            _money(precio_u) if precio_u is not None else "—",
            _money(precio_t) if precio_t is not None else "—",
        ]

        _write_row(
            ws, row, values,
            fonts=[_font(size=10)] * num_cols,
            fills=[fill_row] * num_cols if fill_row else [None] * num_cols,
            aligns=col_aligns,
        )

        if precio_t:
            total_general += Decimal(str(precio_t))
        row += 1

    # ── Fila total ────────────────────────────────────────────────────────
    total_fill = _fill(_SUBTOTAL_BG)
    total_font = _font(bold=True, size=10)
    _merge_write(
        ws, row, 1, num_cols - 1,
        "TOTAL GENERAL",
        font=total_font,
        fill=total_fill,
        align=_align("right"),
        height=22,
    )
    total_cell = ws.cell(row=row, column=num_cols, value=_money(total_general))
    total_cell.font  = _font(bold=True, size=11)
    total_cell.fill  = total_fill
    total_cell.alignment = _align("right")
    total_cell.border = _border()


# ════════════════════════════════════════════════════════════════════════════
# 3. FUNCIONES PÚBLICAS
# ════════════════════════════════════════════════════════════════════════════

def build_despiece_xlsx(proyecto) -> io.BytesIO:
    """
    Genera un xlsx con una hoja 'Despiece' para el proyecto dado.
    Retorna un BytesIO listo para HttpResponse.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Despiece"
    ws.sheet_view.showGridLines = True

    _build_despiece_sheet(ws, proyecto)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_despiece_maestro_xlsx(despiece) -> io.BytesIO:
    """
    Genera un xlsx con una hoja 'Despiece' para un DespieceMaestro.
    Lee DespieceMaestroLinea directamente (modelo ingeniería, no legacy).
    Retorna un BytesIO listo para HttpResponse.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Despiece"
    ws.sheet_view.showGridLines = True

    _build_despiece_maestro_sheet(ws, despiece)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


_INVALID_SHEET_CHARS = str.maketrans({c: "" for c in r"\/?*[]:"})


def _safe_sheet_title(text: str, max_len: int = 28) -> str:
    """
    Sanitiza un string para usarlo como nombre de pestaña Excel.
    Excel prohíbe los caracteres: \ / ? * [ ] :
    Limita la longitud a `max_len` caracteres.
    """
    return text.translate(_INVALID_SHEET_CHARS)[:max_len].strip() or "APU"


def build_apu_xlsx(apu, apus_adicionales=None) -> io.BytesIO:
    """
    Genera un xlsx con:
      - Hoja 'Despiece': listado de materiales del proyecto vinculado al APU.
      - Hoja 'APU <nombre>': análisis de precios unitarios del APU principal.
      - (Opcional) Una hoja adicional por cada APU consolidado en
        `apus_adicionales` (lista de objetos APUProyecto). Ej: cuando hay
        2 APUs juntos → Despiece + APU Principal + APU Adicional = 3 pestañas.

    Retorna un BytesIO listo para HttpResponse.
    """
    wb = Workbook()

    # ── Hoja 1: Despiece ──────────────────────────────────────────────────
    ws_despiece = wb.active
    ws_despiece.title = "Despiece"
    proyecto = None
    if apu.proyecto_sistema_id:
        proyecto = apu.proyecto_sistema.proyecto

    if proyecto:
        _build_despiece_sheet(ws_despiece, proyecto)
    else:
        ws_despiece.cell(row=1, column=1, value="Sin proyecto vinculado")

    # ── Hoja APU principal ────────────────────────────────────────────────
    # Si hay APUs adicionales le ponemos nombre corto al tab para diferenciar.
    if apus_adicionales:
        titulo_principal = _safe_sheet_title(apu.nombre or f"APU {apu.pk}")
        ws_apu = wb.create_sheet(title=titulo_principal)
    else:
        ws_apu = wb.create_sheet(title="APU")
    _build_apu_sheet(ws_apu, apu)

    # ── Hojas APUs adicionales (solo cuando hay consolidación) ────────────
    if apus_adicionales:
        for idx, otro_apu in enumerate(apus_adicionales, start=2):
            titulo_base = _safe_sheet_title(otro_apu.nombre or f"APU {otro_apu.pk}")
            # Evitar nombres de pestaña duplicados
            existing_titles = [s.title for s in wb.worksheets]
            titulo_tab = titulo_base
            if titulo_tab in existing_titles:
                titulo_tab = _safe_sheet_title(
                    (otro_apu.nombre or f"APU {otro_apu.pk}"), max_len=24
                ) + f" ({idx})"
            ws_extra = wb.create_sheet(title=titulo_tab)
            _build_apu_sheet(ws_extra, otro_apu)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
